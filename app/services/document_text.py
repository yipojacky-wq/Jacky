from io import BytesIO
from pathlib import Path

from docx import Document


def extract_upload_text(filename: str, content_type: str, raw: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".docx" or content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx_text(raw)
    if suffix == ".pdf" or content_type == "application/pdf":
        return _extract_pdf_text(raw)
    if _looks_like_zip_or_binary(raw):
        raise ValueError("此檔案看起來是二進位文件。請上傳 .txt、.md、.docx 或可抽取文字的 .pdf。")
    return raw.decode("utf-8-sig", errors="ignore")


def _extract_docx_text(raw: bytes) -> str:
    document = Document(BytesIO(raw))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf_text(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except Exception as exc:
        raise ValueError("缺少 PDF 解析套件，請重新執行 run-windows.bat 安裝 requirements。") from exc

    reader = PdfReader(BytesIO(raw))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[PDF page {index}]\n{text}")
    extracted = "\n\n".join(pages).strip()
    if not extracted:
        raise ValueError("此 PDF 未抽取到文字；若是掃描圖檔，請先 OCR 後再上傳。")
    return extracted


def _looks_like_zip_or_binary(raw: bytes) -> bool:
    if raw.startswith(b"PK\x03\x04"):
        return True
    sample = raw[:2048]
    if not sample:
        return False
    return b"\x00" in sample
